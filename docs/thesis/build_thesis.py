"""Dựng luận văn DOCX (và PDF) từ các chương Markdown trong docs/thesis.

Quy ước Markdown:
  # Chương N. Tiêu đề        -> Heading 1 (sang trang mới)
  ## 1.1. Tiêu đề            -> Heading 2     ### 1.1.1. ...  -> Heading 3
  ![Chú thích](assets/x.png) -> hình, tự đánh số "Hình n" (trường SEQ)
  Dòng "Bảng: chú thích" ngay trước bảng -> tự đánh số "Bảng n"
  ```...```                  -> khối mã;  > ...  -> trích dẫn;  - / 1.  -> danh sách

Chạy: .venv\\Scripts\\python.exe docs/thesis/build_thesis.py [--no-word]
Với Word cài sẵn, script cập nhật mục lục/danh mục hình-bảng và xuất PDF để đếm trang.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
CHAPTERS = [
    "00-tom-tat.md", "01-mo-dau.md", "02-co-so-ly-thuyet.md", "03-phan-tich-yeu-cau.md",
    "04-thiet-ke.md", "05-cai-dat.md", "06-huong-dan-thuc-hanh.md", "07-thuc-nghiem.md",
    "08-thao-luan.md", "09-ket-luan.md", "10-tai-lieu-tham-khao.md", "11-phu-luc.md",
]
BODY_FONT = "Times New Roman"
CODE_FONT = "Consolas"
LINE_SPACING = 1.15

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|(?<![\*\w])\*[^\*\s][^\*]*?\*(?!\w))")


# --------------------------------------------------------------------------- helpers
def set_font(run, name: str = BODY_FONT, size: float | None = None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instr: str, placeholder: str = "", size: float | None = None, bold=None):
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", placeholder), ("end", None)):
        run = paragraph.add_run()
        set_font(run, size=size, bold=bold)
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
            run._r.append(el)
        elif kind == "text":
            run.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)


def shade(element_pr, fill: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    element_pr.append(shd)


def add_inline(paragraph, text: str, size: float | None = None, bold=None, italic=None):
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            set_font(paragraph.add_run(part[2:-2]), size=size, bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            set_font(paragraph.add_run(part[1:-1]), name=CODE_FONT, size=(size or 13) - 2)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            set_font(paragraph.add_run(part[1:-1]), size=size, bold=bold, italic=True)
        else:
            set_font(paragraph.add_run(part), size=size, bold=bold, italic=italic)


def fmt_par(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=None, left=None, after=6, before=0,
            spacing=LINE_SPACING, keep_next=False):
    pf = p.paragraph_format
    p.alignment = align
    pf.space_after, pf.space_before = Pt(after), Pt(before)
    pf.line_spacing = spacing
    if first_indent is not None:
        pf.first_line_indent = Cm(first_indent)
    if left is not None:
        pf.left_indent = Cm(left)
    pf.keep_with_next = keep_next


def setup_styles(doc: Document):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(13)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    sizes = {"Heading 1": (16, True, False), "Heading 2": (14, True, False), "Heading 3": (13, True, True)}
    for name, (size, bold, italic) in sizes.items():
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.italic = BODY_FONT, Pt(size), bold, italic
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), BODY_FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if fonts.get(qn(a)) is not None:
                del fonts.attrib[qn(a)]
        st.paragraph_format.space_before = Pt(18 if name == "Heading 1" else 12)
        st.paragraph_format.space_after = Pt(10 if name == "Heading 1" else 6)
        st.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    cap.font.name, cap.font.size, cap.font.bold, cap.font.italic = BODY_FONT, Pt(11.5), True, False
    cap.font.color.rgb = RGBColor(0, 0, 0)


def page_number_footer(section, fmt: str | None, start: int | None):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(p, "PAGE", "1")
    pg = OxmlElement("w:pgNumType")
    if fmt:
        pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    section._sectPr.append(pg)


# --------------------------------------------------------------------------- front matter
def cover(doc: Document):
    def line(text, size=14, bold=False, italic=False, after=6):
        p = doc.add_paragraph()
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=after, spacing=1.15)
        set_font(p.add_run(text), size=size, bold=bold, italic=italic)

    line("[TÊN TRƯỜNG / KHOA]", 14, True)
    line("ĐỒ ÁN / LUẬN VĂN", 14, True, after=60)
    line("TRỢ LÝ SĂN MỐI ĐE DỌA DỰA TRÊN PEAK ASSISTANT", 22, True, after=12)
    line("Thực thi tất định trên telemetry và khuyến nghị hỗ trợ người quyết định", 15, False, True, after=12)
    line("(AI Agent Hunting v7: tích hợp PEAK Assistant của Cisco Talos, kiểm chứng trên bộ dữ liệu Boss of the SOC v1)",
         12, False, True, after=120)
    line("Sinh viên thực hiện: [Họ và tên]", 13)
    line("Mã số sinh viên: [MSSV]", 13)
    line("Giảng viên hướng dẫn: [Học hàm, học vị, họ và tên]", 13, after=60)
    line("[Địa điểm], [tháng/năm]", 13)


def simple_page(doc: Document, title: str, paragraphs: list[str]):
    doc.add_page_break()
    p = doc.add_paragraph()
    fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    set_font(p.add_run(title), size=16, bold=True)
    for text in paragraphs:
        q = doc.add_paragraph()
        fmt_par(q, first_indent=1.0)
        add_inline(q, text)


def toc_page(doc: Document, title: str, field: str):
    doc.add_page_break()
    p = doc.add_paragraph()
    fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    set_font(p.add_run(title), size=16, bold=True)
    q = doc.add_paragraph()
    fmt_par(q, align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.15)
    add_field(q, field, "(Mục này được Word cập nhật tự động: nhấn Ctrl+A rồi F9)")


# --------------------------------------------------------------------------- body rendering
class Renderer:
    def __init__(self, doc: Document):
        self.doc = doc
        self.pending_table_caption: str | None = None

    def caption(self, kind: str, text: str, *, above: bool):
        p = self.doc.add_paragraph(style="Caption")
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=6 if above else 10, before=6 if above else 2,
                spacing=1.1, keep_next=above)
        set_font(p.add_run(f"{kind} "), size=11.5, bold=True)
        add_field(p, f"SEQ {kind} \\* ARABIC", "0", size=11.5, bold=True)
        set_font(p.add_run(". "), size=11.5, bold=True)
        add_inline(p, text, size=11.5, bold=False, italic=True)

    def heading(self, level: int, text: str):
        h = self.doc.add_heading(text, level=min(level, 3))
        if level == 1:
            h.paragraph_format.page_break_before = True
        for r in h.runs:
            set_font(r, size={1: 16, 2: 14, 3: 13}[min(level, 3)], bold=True, italic=(level >= 3))
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT

    def code(self, lines: list[str]):
        for i, ln in enumerate(lines):
            p = self.doc.add_paragraph()
            fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, left=0.4, after=0, spacing=1.0)
            shade(p._p.get_or_add_pPr(), "F2F2F2")
            if i == 0:
                p.paragraph_format.space_before = Pt(4)
            if i == len(lines) - 1:
                p.paragraph_format.space_after = Pt(8)
            set_font(p.add_run(ln if ln else " "), name=CODE_FONT, size=9)

    def table(self, header: list[str], rows: list[list[str]]):
        if self.pending_table_caption:
            self.caption("Bảng", self.pending_table_caption, above=True)
            self.pending_table_caption = None
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(header):
            c = t.rows[0].cells[j]
            c.paragraphs[0].text = ""
            fmt_par(c.paragraphs[0], align=WD_ALIGN_PARAGRAPH.LEFT, after=2, before=2, spacing=1.0)
            add_inline(c.paragraphs[0], h, size=10.5, bold=True)
            shade(c._tc.get_or_add_tcPr(), "D9E2F3")
        for r in rows:
            cells = t.add_row().cells
            for j in range(len(header)):
                cells[j].paragraphs[0].text = ""
                fmt_par(cells[j].paragraphs[0], align=WD_ALIGN_PARAGRAPH.LEFT, after=2, before=2, spacing=1.0)
                add_inline(cells[j].paragraphs[0], r[j] if j < len(r) else "", size=10.5)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader")
        el.set(qn("w:val"), "true")
        trpr.append(el)
        spacer = self.doc.add_paragraph()
        fmt_par(spacer, after=4, spacing=1.0)

    def figure(self, caption: str, path: Path):
        p = self.doc.add_paragraph()
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, before=6, spacing=1.0, keep_next=True)
        p.add_run().add_picture(str(path), width=Cm(15.5))
        self.caption("Hình", caption, above=False)

    def render(self, text: str):
        lines = text.splitlines()
        i, n = 0, len(lines)
        while i < n:
            line = lines[i]
            if line.strip().startswith("```"):
                block: list[str] = []
                i += 1
                while i < n and not lines[i].strip().startswith("```"):
                    block.append(lines[i].rstrip("\n"))
                    i += 1
                i += 1
                self.code(block)
                continue
            m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", line.strip())
            if m:
                self.figure(m.group(1), HERE / m.group(2))
                i += 1
                continue
            if line.lstrip().startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
                header = [c.strip() for c in line.strip().strip("|").split("|")]
                i += 2
                rows = []
                while i < n and lines[i].lstrip().startswith("|"):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                    i += 1
                self.table(header, rows)
                continue
            cap = re.match(r"^Bảng:\s*(.+)$", line.strip())
            if cap:
                self.pending_table_caption = cap.group(1)
                i += 1
                continue
            h = re.match(r"^(#{1,3})\s+(.*)$", line)
            if h:
                self.heading(len(h.group(1)), h.group(2).strip())
                i += 1
                continue
            if line.lstrip().startswith(">"):
                p = self.doc.add_paragraph()
                fmt_par(p, left=1.0, after=6)
                add_inline(p, line.lstrip().lstrip(">").strip(), italic=True)
                i += 1
                continue
            b = re.match(r"^(\s*)[-*]\s+(.*)$", line)
            if b:
                p = self.doc.add_paragraph()
                depth = len(b.group(1)) // 2
                fmt_par(p, left=1.0 + 0.8 * depth, first_indent=-0.5, after=3)
                set_font(p.add_run("•\t" if depth == 0 else "–\t"))
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0 + 0.8 * depth))
                add_inline(p, b.group(2))
                i += 1
                continue
            ref = re.match(r"^\[(\d+)\]\s+(.*)$", line.strip())
            if ref:
                p = self.doc.add_paragraph()
                fmt_par(p, align=WD_ALIGN_PARAGRAPH.LEFT, left=1.0, first_indent=-1.0, after=4, spacing=1.15)
                set_font(p.add_run(f"[{ref.group(1)}]\t"))
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
                add_inline(p, ref.group(2))
                i += 1
                continue
            nm = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
            if nm:
                p = self.doc.add_paragraph()
                fmt_par(p, left=1.0, first_indent=-0.7, after=3)
                set_font(p.add_run(f"{nm.group(1)}.\t"))
                p.paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
                add_inline(p, nm.group(2))
                i += 1
                continue
            if line.strip() in ("---", "***", "___") or not line.strip():
                i += 1
                continue
            p = self.doc.add_paragraph()
            fmt_par(p, first_indent=1.0)
            add_inline(p, line.strip())
            i += 1


# --------------------------------------------------------------------------- main
def build(out: Path) -> None:
    doc = Document()
    setup_styles(doc)
    cover(doc)

    front = doc.add_section(WD_SECTION.NEW_PAGE)
    page_number_footer(front, "lowerRoman", 1)
    simple_page(doc, "LỜI CAM ĐOAN", [
        "Tôi xin cam đoan đây là công trình nghiên cứu của riêng tôi dưới sự hướng dẫn của giảng viên hướng dẫn. "
        "Các số liệu và kết quả thực nghiệm trong đồ án là trung thực, được sinh ra từ mã nguồn và dữ liệu đi kèm "
        "(kho `AI_Agent_Hunting`, thư mục `results/`); các tài liệu tham khảo đều được trích dẫn rõ nguồn.",
        "Họ tên sinh viên: [Họ và tên]          Chữ ký: ____________________",
    ])

    parts = {name: (HERE / name).read_text(encoding="utf-8") for name in CHAPTERS if (HERE / name).exists()}
    r = Renderer(doc)
    r.render(parts.pop("00-tom-tat.md", ""))
    toc_page(doc, "MỤC LỤC", 'TOC \\o "1-2" \\h \\z \\u')
    toc_page(doc, "DANH MỤC HÌNH", 'TOC \\h \\z \\c "Hình"')
    toc_page(doc, "DANH MỤC BẢNG", 'TOC \\h \\z \\c "Bảng"')

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    page_number_footer(body, "decimal", 1)
    for name in CHAPTERS:
        if name in parts:
            r.render(parts[name])
    doc.save(str(out))
    print(f"[+] DOCX: {out}")


def word_update(docx: Path, pdf: Path) -> int | None:
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print("[!] pywin32 không có; bỏ qua cập nhật mục lục")
        return None
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    try:
        d = word.Documents.Open(str(docx.resolve()))
        d.Repaginate()
        d.Fields.Update()
        for _ in range(2):
            for toc in d.TablesOfContents:
                toc.Update()
            for toc in d.TablesOfFigures:
                toc.Update()
            d.Repaginate()
        pages = d.ComputeStatistics(2)
        d.Save()
        d.ExportAsFixedFormat(str(pdf.resolve()), 17)
        d.Close(False)
        return int(pages)
    finally:
        word.Quit()


def main() -> int:
    out = HERE / "LUAN-VAN-AI-AGENT-HUNTING.docx"
    build(out)
    if "--no-word" not in sys.argv:
        pages = word_update(out, out.with_suffix(".pdf"))
        if pages:
            print(f"[+] Tổng số trang (gồm bìa, mục lục, phụ lục): {pages}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
